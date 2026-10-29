"""File -> plain text. Only the extracted text is ever stored; nothing is executed."""
import io
import zipfile
from pathlib import PurePosixPath

ALLOWED_EXTENSIONS = {".txt", ".md", ".docx", ".pdf"}
MAX_DOCX_UNCOMPRESSED = 50 * 1024 * 1024
MAX_PDF_PAGES = 300


class IngestError(Exception):
    def __init__(self, status_code: int, message: str):
        super().__init__(message)
        self.status_code = status_code
        self.message = message


def extension_of(filename: str) -> str:
    return PurePosixPath(filename.lower()).suffix


def clean_text(text: str) -> str:
    # Postgres TEXT cannot contain NUL bytes.
    return text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n").strip()


def _extract_plain(data: bytes) -> str:
    if b"\x00" in data:
        raise IngestError(422, "The file does not look like a text file (binary content).")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise IngestError(422, "Text files must be UTF-8 encoded.")


def _extract_docx(data: bytes) -> str:
    if not data.startswith(b"PK"):
        raise IngestError(422, "The .docx file is corrupted (not a valid Word document).")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(i.file_size for i in z.infolist()) > MAX_DOCX_UNCOMPRESSED:
                raise IngestError(413, "The .docx file expands to too much data.")
            if "word/document.xml" not in z.namelist():
                raise IngestError(422, "The .docx file is corrupted (no document body).")
        from docx import Document

        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(parts)
    except IngestError:
        raise
    except Exception:
        raise IngestError(422, "The .docx file is corrupted or unreadable.")


def _extract_pdf(data: bytes) -> str:
    if b"%PDF" not in data[:1024]:
        raise IngestError(422, "The PDF file is corrupted (missing PDF header).")
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise IngestError(422, "Password-protected PDFs are not supported.")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise IngestError(413, f"PDF has too many pages (max {MAX_PDF_PAGES}).")
        text = "\n\n".join((page.extract_text() or "") for page in reader.pages)
    except IngestError:
        raise
    except Exception:
        raise IngestError(422, "The PDF file is corrupted or unreadable.")
    if not text.strip():
        raise IngestError(
            422,
            "The PDF has no extractable text (empty or scanned). OCR is not supported.",
        )
    return text


def extract_text(filename: str, data: bytes) -> str:
    ext = extension_of(filename)
    if ext not in ALLOWED_EXTENSIONS:
        raise IngestError(
            415,
            f"Unsupported file type '{ext or 'none'}'. Allowed: .txt, .md, .docx, .pdf",
        )
    if not data:
        raise IngestError(422, "The file is empty.")
    extractor = {
        ".txt": _extract_plain,
        ".md": _extract_plain,
        ".docx": _extract_docx,
        ".pdf": _extract_pdf,
    }[ext]
    text = clean_text(extractor(data))
    if not text:
        raise IngestError(422, "The file contains no text.")
    return text
